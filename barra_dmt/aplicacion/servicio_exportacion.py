"""Exportacion de reportes a Excel (.xlsx).

Dos reportes, pensados para descargarse desde la pestaña Caja:
- reporte de una caja (turno) puntual: cabecera + ventas + top productos.
- reporte historico: analisis agregado de todas las cajas + pronostico.

Usa openpyxl directo (sin pandas) para no engordar el ejecutable de PyInstaller
mas de lo necesario -- misma logica que ya sigue ServicioPronostico.
"""
from __future__ import annotations

import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..dominio.entidades import CierreCaja, SesionCaja
from ..dominio.repositorios import RepositorioVentas
from .servicio_analisis import ServicioAnalisis
from .servicio_caja import ServicioCaja, TotalesSesion
from .servicio_pronostico import ServicioPronostico

_FUENTE_TITULO = Font(bold=True, size=13)
_FUENTE_ENCABEZADO = Font(bold=True, color="FFFFFF")
_RELLENO_ENCABEZADO = PatternFill("solid", fgColor="2F6F4F")


def _fecha(ts: Optional[float]) -> str:
    return datetime.datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M") if ts else "-"


def _titulo(hoja: Worksheet, texto: str, fila: int = 1) -> int:
    hoja.cell(row=fila, column=1, value=texto).font = _FUENTE_TITULO
    return fila + 2


def _tabla(hoja: Worksheet, fila: int, encabezados: list[str], filas: list[list]) -> int:
    for col, texto in enumerate(encabezados, start=1):
        celda = hoja.cell(row=fila, column=col, value=texto)
        celda.font = _FUENTE_ENCABEZADO
        celda.fill = _RELLENO_ENCABEZADO
    for i, valores in enumerate(filas, start=1):
        for col, valor in enumerate(valores, start=1):
            hoja.cell(row=fila + i, column=col, value=valor)
    for col in range(1, len(encabezados) + 1):
        hoja.column_dimensions[get_column_letter(col)].width = 22
    return fila + len(filas) + 2


class ServicioExportacion:
    def __init__(self, servicio_caja: ServicioCaja, servicio_analisis: ServicioAnalisis,
                 servicio_pronostico: ServicioPronostico, repo_ventas: RepositorioVentas):
        self._servicio_caja = servicio_caja
        self._servicio_analisis = servicio_analisis
        self._servicio_pronostico = servicio_pronostico
        self._repo_ventas = repo_ventas

    # ---------------------------------------------------------- reporte de una caja
    def exportar_reporte_caja_abierta(self, ruta: Path, sesion: SesionCaja, totales: TotalesSesion) -> None:
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Resumen"
        fila = _titulo(hoja, f"Caja abierta - {sesion.cajero}")
        _tabla(hoja, fila, ["Campo", "Valor"], [
            ["Cajero", sesion.cajero],
            ["Abierta", _fecha(sesion.abierta_en)],
            ["Estado", "Turno en curso"],
            ["Efectivo inicial", sesion.efectivo_inicial],
            ["Ventas del turno", totales.cantidad_ventas],
            ["Total vendido", totales.total_ventas],
            ["Total efectivo", totales.total_efectivo],
            ["Total QR", totales.total_qr],
            ["Gastado en DMT (costo)", totales.total_gastos_dmt],
            ["Efectivo esperado", totales.efectivo_esperado],
        ])
        self._agregar_ventas_y_top(libro, sesion.id)
        libro.save(ruta)

    def exportar_reporte_caja_cerrada(self, ruta: Path, cierre: CierreCaja) -> None:
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Resumen"
        fila = _titulo(hoja, f"Caja cerrada - {cierre.cajero}")
        _tabla(hoja, fila, ["Campo", "Valor"], [
            ["Cajero", cierre.cajero],
            ["Abierta", _fecha(cierre.abierta_en)],
            ["Cerrada", _fecha(cierre.cerrada_en)],
            ["Efectivo inicial", cierre.efectivo_inicial],
            ["Total vendido", cierre.total_ventas],
            ["Total efectivo", cierre.total_efectivo],
            ["Total QR", cierre.total_qr],
            ["Gastado en DMT (costo)", cierre.total_gastos_dmt],
            ["Efectivo esperado", cierre.efectivo_esperado],
            ["Efectivo contado", cierre.efectivo_contado],
            ["Diferencia", cierre.diferencia],
        ])
        self._agregar_ventas_y_top(libro, cierre.sesion_id)
        libro.save(ruta)

    def _agregar_ventas_y_top(self, libro: Workbook, sesion_id: str) -> None:
        ventas = sorted((v for v in self._repo_ventas.listar_por_sesion(sesion_id) if not v.anulada),
                         key=lambda v: v.timestamp)
        hoja_ventas = libro.create_sheet("Ventas")
        fila = _titulo(hoja_ventas, "Ventas del turno")
        _tabla(hoja_ventas, fila, ["Hora", "Metodo de pago", "Total", "Items"], [
            [_fecha(v.timestamp), v.metodo_pago, v.total,
             ", ".join(f"{it.cantidad}x {it.nombre}" for it in v.items)]
            for v in ventas
        ])

        top = self._servicio_analisis.top_productos(sesion_id=sesion_id, n=50)
        hoja_top = libro.create_sheet("Top productos")
        fila = _titulo(hoja_top, "Productos mas vendidos en el turno")
        _tabla(hoja_top, fila, ["Producto", "Unidades", "Ingresos"],
               [[f["nombre"], f["cantidad"], f["ingresos"]] for f in top])

    # ---------------------------------------------------------- reporte historico
    def exportar_reporte_historico(self, ruta: Path) -> None:
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Resumen"
        resumen = self._servicio_analisis.resumen()
        estadisticas = self._servicio_analisis.estadisticas()
        fila = _titulo(hoja, "Resumen historico")
        _tabla(hoja, fila, ["Metrica", "Valor"], [
            ["Total vendido", resumen["total_vendido"]],
            ["Cantidad de ventas", resumen["n_ventas"]],
            ["Unidades vendidas", resumen["unidades"]],
            ["Ticket promedio", round(resumen["ticket_promedio"])],
            ["Utilidad estimada", resumen["utilidad"]],
            ["Gastado en DMT (costo)", resumen["gasto_dmt"]],
            ["Ticket minimo", estadisticas["minimo"]],
            ["Ticket maximo", estadisticas["maximo"]],
            ["Ticket mediana", round(estadisticas["mediana"])],
            ["Desviacion estandar", round(estadisticas["desviacion"])],
        ])

        hoja_cat = libro.create_sheet("Por categoria")
        fila = _titulo(hoja_cat, "Ventas por categoria")
        _tabla(hoja_cat, fila, ["Categoria", "Total"],
               [[f["categoria"], f["total"]] for f in self._servicio_analisis.ventas_por_categoria()])

        hoja_pago = libro.create_sheet("Por metodo de pago")
        fila = _titulo(hoja_pago, "Ventas por metodo de pago")
        _tabla(hoja_pago, fila, ["Metodo", "Total"],
               [[m, t] for m, t in self._servicio_analisis.ventas_por_metodo_pago().items()])

        hoja_top = libro.create_sheet("Top productos")
        fila = _titulo(hoja_top, "Productos mas vendidos (historico)")
        _tabla(hoja_top, fila, ["Producto", "Unidades", "Ingresos"],
               [[f["nombre"], f["cantidad"], f["ingresos"]]
                for f in self._servicio_analisis.top_productos(n=50)])

        hoja_hora = libro.create_sheet("Por hora")
        fila = _titulo(hoja_hora, "Ventas por hora del dia")
        _tabla(hoja_hora, fila, ["Hora", "Total"],
               [[f"{h:02d}:00", t] for h, t in enumerate(self._servicio_analisis.ventas_por_hora())])

        hoja_stock = libro.create_sheet("Stock bajo")
        fila = _titulo(hoja_stock, "Productos con stock bajo (actual)")
        _tabla(hoja_stock, fila, ["Producto", "Stock", "Stock minimo"],
               [[p.nombre, p.stock, p.stock_minimo] for p in self._servicio_analisis.productos_stock_bajo()])

        hoja_gastos = libro.create_sheet("Gastos DMT por producto")
        fila = _titulo(hoja_gastos, "Productos entregados como DMT (costo)")
        gastos_por_producto = sorted(
            self._servicio_analisis.gasto_dmt_por_producto().values(),
            key=lambda f: f["valor"], reverse=True)
        _tabla(hoja_gastos, fila, ["Producto", "Unidades", "Valor (costo)"],
               [[f["nombre"], f["cantidad"], f["valor"]] for f in gastos_por_producto])

        hoja_repo = libro.create_sheet("Reposiciones")
        fila = _titulo(hoja_repo, "Reposiciones de stock por producto")
        _tabla(hoja_repo, fila, ["Producto", "Unidades repuestas"],
               [[f["nombre"], f["cantidad"]] for f in self._servicio_analisis.total_repuesto_por_producto()])

        hoja_turnos = libro.create_sheet("Turnos (cajas)")
        fila = _titulo(hoja_turnos, "Historial de turnos cerrados")
        cierres = sorted(self._servicio_caja.listar_cierres(), key=lambda c: c.cerrada_en)
        _tabla(hoja_turnos, fila,
               ["Cajero", "Abierta", "Cerrada", "Total vendido", "Efectivo esperado",
                "Efectivo contado", "Diferencia"],
               [[c.cajero, _fecha(c.abierta_en), _fecha(c.cerrada_en), c.total_ventas,
                 c.efectivo_esperado, c.efectivo_contado, c.diferencia] for c in cierres])

        hoja_margen = libro.create_sheet("Margen por producto")
        fila = _titulo(hoja_margen, "Pronostico: margen por producto")
        _tabla(hoja_margen, fila, ["Producto", "Categoria", "Precio", "Costo", "Margen", "Margen %"],
               [[m["nombre"], m["categoria"], m["precio"], m["costo"], m["margen_abs"], round(m["margen_pct"], 1)]
                for m in self._servicio_pronostico.margen_por_producto()])

        hoja_reco = libro.create_sheet("Pronostico proximo evento")
        fila = _titulo(hoja_reco, "Stock recomendado para el proximo evento")
        _tabla(hoja_reco, fila,
               ["Producto", "Stock actual", "Recomendado (unid)", "Recomendado (pacas)",
                "Fuente", "Eventos considerados"],
               [[r.nombre, r.stock_actual, r.recomendado_unidades, r.recomendado_pacas,
                 r.fuente, r.eventos_considerados]
                for r in self._servicio_pronostico.pronostico_proximo_evento()])

        hoja_proy = libro.create_sheet("Proyeccion ventas")
        proyeccion = self._servicio_pronostico.proyeccion_ventas(granularidad="dia", horizonte=30)
        fila = _titulo(hoja_proy, "Proyeccion de ventas (historico + proximos 30 dias)")
        filas = [[p.etiqueta, round(p.total), "historico"] for p in proyeccion.historico]
        filas += [[p.etiqueta, round(p.total), "proyectado"] for p in proyeccion.proyectado]
        _tabla(hoja_proy, fila, ["Periodo", "Total", "Tipo"], filas)

        libro.save(ruta)
